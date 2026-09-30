"""
KI-Sparringspartner für den Kurs Entrepreneurship
Oberfläche: Streamlit Community Cloud (kostenlos)
Sprachmodell: Apertus über die Public AI Inference Utility (Free Plan)
"""

import io
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import streamlit as st
from openai import OpenAI
from pypdf import PdfReader
from docx import Document

# ---------------------------------------------------------------------------
# Einstellungen
# ---------------------------------------------------------------------------
API_BASE = "https://api.publicai.co/v1"
MODEL = st.secrets.get("MODEL", "swiss-ai/apertus-v1.5-8b")
MAX_CANVAS_CHARS = 15000      # Canvas-Text wird bei Bedarf gekürzt
MAX_HISTORY = 16              # nur die letzten Nachrichten werden mitgeschickt

# Kontingent-Anzeige (Werte in US-Dollar, in den Secrets einstellbar)
BUDGET_START = float(st.secrets.get("BUDGET_START", 2.00))
BUDGET_REMAINING = float(st.secrets.get("BUDGET_REMAINING", BUDGET_START))
PRICE_IN = float(st.secrets.get("PRICE_IN", 0.10))    # pro 1 Mio. Tokens
PRICE_OUT = float(st.secrets.get("PRICE_OUT", 0.20))  # pro 1 Mio. Tokens

SYSTEM_PROMPT = """Du bist ein erfahrener, kritischer und wohlwollender Sparringspartner \
für Studierende im Kurs Entrepreneurship an einer Schweizer Hochschule. \
Du schreibst auf Deutsch in Schweizer Rechtschreibung (ss statt ß).

Deine Rolle:
- Du hilfst den Studierenden, ihre Geschäftsidee und ihr Business Model Canvas \
(oder Lean Canvas) zu schärfen. Die neun Bausteine des Business Model Canvas sind \
Kundensegmente, Wertangebot, Kanäle, Kundenbeziehungen, Einnahmequellen, \
Schlüsselressourcen, Schlüsselaktivitäten, Schlüsselpartner und Kostenstruktur.
- Du bist ein Sparringspartner, kein Ghostwriter. Du schreibst den Studierenden \
kein fertiges Canvas und keinen Businessplan. Du stellst gezielte Fragen, \
deckst Annahmen auf und forderst Belege.
- Du prüfst besonders die Konsistenz zwischen den Bausteinen, zum Beispiel ob \
das Wertangebot wirklich zum Kundensegment passt und ob die Einnahmen die Kosten decken.
- Du benennst Schwächen offen, aber respektvoll, und würdigst auch Stärken.
- Du regst an, riskante Annahmen mit einfachen Experimenten zu testen \
(Kundeninterviews, Landingpage, Prototyp).

Stil:
- Kurz und konkret. Höchstens drei Kritikpunkte oder Fragen pro Antwort, \
damit die Studierenden nicht überfordert werden.
- Ende jede Antwort mit genau einer Rückfrage, die zum Weiterdenken anregt.
- Wenn du etwas nicht weisst oder Marktzahlen nicht belegen kannst, sag das ehrlich \
und erfinde keine Zahlen oder Quellen.
- Bleib beim Thema Entrepreneurship. Bei fachfremden Anfragen weist du freundlich \
auf deine Rolle hin."""

WELCOME = (
    "Hallo! Ich bin dein Sparringspartner für deine Geschäftsidee. "
    "Lade links dein Canvas hoch oder beschreib mir deine Idee in ein paar Sätzen. "
    "Womit möchtest du anfangen?"
)

APP_NAME = st.secrets.get("APP_NAME", "KI-Sparringspartner Entrepreneurship")
st.set_page_config(page_title=APP_NAME, page_icon="💡")

# ---------------------------------------------------------------------------
# Logo (optional): Datei "logo.png" neben app.py auf GitHub hochladen
# ---------------------------------------------------------------------------
LOGO = Path(__file__).parent / "logo.png"
if LOGO.exists():
    st.logo(str(LOGO), size="large")

# ---------------------------------------------------------------------------
# Sicherheitshinweise
# ---------------------------------------------------------------------------
HINWEISE = """
- Gib keine persönlichen Daten ein (Namen, Adressen, Telefonnummern, Gesundheitsdaten).
- Lade keine vertraulichen Unterlagen hoch, etwa Verträge, Finanzdaten Dritter oder \
Geschäftsgeheimnisse von Partnerfirmen.
- Die KI kann sich irren und Fakten erfinden. Prüfe Zahlen, Marktdaten und Quellen selbst.
- Die Antworten sind Denkanstösse, keine Bewertung und keine Rechts-, Steuer- oder Finanzberatung.
- Deine Eingaben werden zur Beantwortung an den Dienst Public AI übermittelt. \
Die Dozierenden sehen deine Gespräche nicht.
- Die Regeln der Hochschule zu KI in Leistungsnachweisen gelten auch hier.
"""

# ---------------------------------------------------------------------------
# Zugang und Öffnungszeiten (werden in den Streamlit-Secrets eingestellt)
# ---------------------------------------------------------------------------
TAGE = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]


OPEN_DAYS = [t.strip() for t in st.secrets.get("OPEN_DAYS", "Mo,Di,Mi,Do,Fr,Sa,So").split(",")]
OPEN_FROM = st.secrets.get("OPEN_FROM", "00:00")
OPEN_UNTIL = st.secrets.get("OPEN_UNTIL", "23:59")
END_DATE = st.secrets.get("END_DATE", "")  # Format TT.MM.JJJJ, leer = unbegrenzt


def zeiten_text() -> str:
    text = f"Geöffnet {', '.join(OPEN_DAYS)} von {OPEN_FROM} bis {OPEN_UNTIL} Uhr."
    if END_DATE:
        text += f" Verfügbar bis {END_DATE}."
    return text


def zugang_offen() -> tuple[bool, str]:
    if not st.secrets.get("ACCESS_OPEN", True):
        return False, "Der Zugang ist zurzeit geschlossen."
    jetzt = datetime.now(ZoneInfo("Europe/Zurich"))
    if END_DATE:
        ende = datetime.strptime(END_DATE, "%d.%m.%Y").date()
        if jetzt.date() > ende:
            return False, f"Das Angebot war bis {END_DATE} verfügbar und ist nun beendet."
    heute = TAGE[jetzt.weekday()]
    uhrzeit = jetzt.strftime("%H:%M")
    if heute in OPEN_DAYS and OPEN_FROM <= uhrzeit <= OPEN_UNTIL:
        return True, ""
    return False, zeiten_text() + " Schau gerne zu diesen Zeiten wieder vorbei."


offen, meldung = zugang_offen()
if not offen:
    st.title(APP_NAME)
    st.info(meldung)
    st.stop()


# ---------------------------------------------------------------------------
# Passwortschutz
# ---------------------------------------------------------------------------
def check_password() -> bool:
    if st.session_state.get("auth_ok"):
        return True
    st.title(APP_NAME)
    st.write("Bitte gib das Kurspasswort ein, das du von deiner Dozentin oder deinem Dozenten erhalten hast.")
    st.caption("🕒 " + zeiten_text())
    with st.expander("Wichtige Hinweise vor der Nutzung", expanded=True):
        st.markdown(HINWEISE)
    pw = st.text_input("Kurspasswort", type="password")
    if st.button("Anmelden"):
        if pw == st.secrets.get("APP_PASSWORD", ""):
            st.session_state.auth_ok = True
            st.rerun()
        else:
            st.error("Das Passwort stimmt nicht. Prüfe Gross- und Kleinschreibung.")
    return False


if not check_password():
    st.stop()


# ---------------------------------------------------------------------------
# Hilfsfunktionen
# ---------------------------------------------------------------------------
@st.cache_resource
def get_client() -> OpenAI:
    return OpenAI(
        api_key=st.secrets["PUBLICAI_API_KEY"],
        base_url=API_BASE,
        default_headers={"User-Agent": "Entrepreneurship-Sparring/1.0"},
    )


def extract_text(uploaded) -> str:
    data = uploaded.getvalue()
    name = uploaded.name.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    return data.decode("utf-8", errors="ignore")


@st.cache_resource
def verbrauch() -> dict:
    """Geschätzter Verbrauch seit dem letzten Start der App (für alle Nutzenden gemeinsam)."""
    return {"dollar": 0.0, "stand": BUDGET_REMAINING}


def verbrauch_erfassen(messages: list, antwort: str):
    v = verbrauch()
    if v["stand"] != BUDGET_REMAINING:  # Wert in den Secrets wurde aktualisiert
        v["dollar"], v["stand"] = 0.0, BUDGET_REMAINING
    tokens_in = sum(len(m["content"]) for m in messages) / 4
    tokens_out = len(antwort) / 4
    v["dollar"] += tokens_in * PRICE_IN / 1e6 + tokens_out * PRICE_OUT / 1e6


def kontingent_anteil() -> float:
    rest = BUDGET_REMAINING - verbrauch()["dollar"]
    return max(0.0, min(1.0, rest / BUDGET_START)) if BUDGET_START > 0 else 0.0


def build_messages() -> list:
    system = SYSTEM_PROMPT
    canvas = st.session_state.get("canvas_text")
    if canvas:
        system += (
            "\n\nDie studierende Person hat folgendes Canvas hochgeladen. "
            "Beziehe dich darauf:\n<canvas>\n" + canvas + "\n</canvas>"
        )
    history = st.session_state.messages[-MAX_HISTORY:]
    return [{"role": "system", "content": system}] + history


def ask_model():
    stream = get_client().chat.completions.create(
        model=MODEL,
        messages=build_messages(),
        temperature=0.6,
        max_tokens=900,
        stream=True,
    )
    for chunk in stream:
        if chunk.choices and chunk.choices[0].delta.content:
            yield chunk.choices[0].delta.content


def respond(user_text: str):
    st.session_state.messages.append({"role": "user", "content": user_text})
    with st.chat_message("user"):
        st.markdown(user_text)
    with st.chat_message("assistant"):
        try:
            answer = st.write_stream(ask_model())
            verbrauch_erfassen(build_messages(), answer)
        except Exception as e:  # Limits, Netzwerk, Schlüssel
            msg = str(e).lower()
            if any(w in msg for w in ("402", "credit", "balance", "insufficient", "quota", "payment")):
                answer = "Das Kontingent für dieses Semester ist aufgebraucht. Der Sparringspartner steht deshalb zurzeit nicht zur Verfügung."
            elif "429" in msg or "rate" in msg:
                answer = "Gerade sind sehr viele Anfragen gleichzeitig unterwegs. Warte eine Minute und versuche es erneut."
            elif "401" in msg or "auth" in msg:
                answer = "Der Zugang zum Sprachmodell funktioniert nicht. Bitte informiere deine Dozentin oder deinen Dozenten."
            else:
                answer = "Die Antwort konnte nicht erstellt werden. Versuche es in einem Moment noch einmal."
            st.warning(answer)
    st.session_state.messages.append({"role": "assistant", "content": answer})


# ---------------------------------------------------------------------------
# Seitenleiste: Canvas hochladen
# ---------------------------------------------------------------------------
if "messages" not in st.session_state:
    st.session_state.messages = [{"role": "assistant", "content": WELCOME}]

with st.sidebar:
    st.header("Dein Canvas")
    st.caption("PDF, Word (.docx) oder Textdatei. Fotos und eingescannte PDFs kann das Modell nicht lesen.")
    uploaded = st.file_uploader("Datei auswählen", type=["pdf", "docx", "txt", "md"])

    if uploaded is not None and st.session_state.get("canvas_name") != uploaded.name:
        text = extract_text(uploaded).strip()
        if len(text) < 30:
            st.error("In dieser Datei wurde kaum Text gefunden. Exportiere dein Canvas als PDF mit echtem Text oder als Word-Datei.")
        else:
            st.session_state.canvas_text = text[:MAX_CANVAS_CHARS]
            st.session_state.canvas_name = uploaded.name
            if len(text) > MAX_CANVAS_CHARS:
                st.info("Die Datei ist lang, deshalb wird nur der erste Teil berücksichtigt.")

    if st.session_state.get("canvas_text"):
        st.success(f"Geladen: {st.session_state.canvas_name}")
        if st.button("Feedback zu meinem Canvas", type="primary", use_container_width=True):
            st.session_state.pending = (
                "Bitte gib mir ein kritisches Feedback zu meinem hochgeladenen Canvas. "
                "Wo siehst du die grössten Schwachstellen und ungeprüften Annahmen?"
            )

    st.divider()
    if st.button("Neues Gespräch", use_container_width=True):
        for key in ("messages", "canvas_text", "canvas_name", "pending"):
            st.session_state.pop(key, None)
        st.rerun()
    anteil = kontingent_anteil()
    st.progress(anteil, text=f"Kontingent: ca. {round(anteil * 100)} % verfügbar")
    st.caption("🕒 " + zeiten_text())
    with st.expander("Hinweise zur Nutzung"):
        st.markdown(HINWEISE)

# ---------------------------------------------------------------------------
# Chat
# ---------------------------------------------------------------------------
st.title(APP_NAME)

for m in st.session_state.messages:
    with st.chat_message(m["role"]):
        st.markdown(m["content"])

pending = st.session_state.pop("pending", None)
prompt = st.chat_input("Stell eine Frage oder beschreib deine Idee …")
if pending:
    respond(pending)
elif prompt:
    respond(prompt)
